import os
import sys
import unicodedata
from unittest.mock import patch

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from jinja2.exceptions import TemplateSyntaxError

from app.services.cv_generator_fallback import (
    _extract_all_text,
    _build_apilayer_template_pairs,
    _plan_template_execution,
    _supports_fast_placeholder_render,
    process_cv,
)


def test_fallback_uses_placeholder_mode_for_jinja_templates(tmp_path):
    template_path = tmp_path / "placeholder_template.docx"
    doc = Document()
    doc.add_paragraph("{{ full_name }}")
    doc.add_paragraph("{{email}} • {{phone}}")
    doc.add_paragraph("{{ current_position }}")
    doc.add_paragraph("Professional Summary")
    doc.add_paragraph("{{ summary }}")
    doc.add_paragraph("Education")
    doc.add_paragraph("{% for edu in education %}")
    doc.add_paragraph("{{ edu.degree }}")
    doc.add_paragraph("{{ edu.school }}")
    doc.add_paragraph("{% endfor %}")
    doc.add_paragraph("Experience")
    doc.add_paragraph("{% for exp in experience %}")
    doc.add_paragraph("{{ exp.title }}")
    doc.add_paragraph("{{ exp.company }}")
    doc.add_paragraph("{{ exp.period }}")
    doc.add_paragraph("{% endfor %}")
    doc.save(str(template_path))

    employee = {
        "name": "Anouar ABDALLAH",
        "email": "user1@nextrh.com",
        "phone": "+216 27 920 721",
        "title": "Security Consultant",
        "linkedin": "linkedin.com/in/anouar-abdallah",
        "address": "Tunis, Tunisia",
        "summary": "Results-driven security professional.",
        "skills": ["Security", "Networking", "Firewall"],
        "experience": [
            {
                "title": "Security Consultant",
                "company": "Next Step IT",
                "dates": "2020 - Present",
                "description": "Managed SOC operations.",
            }
        ],
        "education": [
            {
                "degree": "BSc Computer Science",
                "institution": "Esprit",
                "dates": "2015 - 2019",
            }
        ],
        "projects": [
            {
                "name": "Cloud Hardening Program",
                "description": "Led hardening for enterprise cloud workloads.",
                "startDate": "2026",
                "endDate": "2026",
            }
        ],
        "certifications": [
            {"name": "CCNP Security", "year": "2020"},
            {"name": "Fortinet NSE 7", "year": "2022"},
        ],
    }

    out = process_cv(
        template_path=str(template_path),
        employee_data=employee,
        output_dir=str(tmp_path),
        output_pdf=False,
    )
    assert os.path.exists(out)

    text, _paras = _extract_all_text(out)
    lower = text.lower()

    assert "anouar abdallah" in lower
    assert "user1@nextrh.com" in lower
    assert "+216 27 920 721" in text
    assert "next step it" in lower
    assert "bsc computer science" in lower
    assert "{{" not in text
    assert "{%" not in text


def test_structural_docxtpl_tags_disable_fast_placeholder_render(tmp_path):
    template_path = tmp_path / "structural_placeholder_template.docx"
    doc = Document()
    doc.add_paragraph("{{ full_name }}")
    doc.add_paragraph("{%tr for exp in experience %}")
    doc.add_paragraph("{{ exp.title }}")
    doc.add_paragraph("{%tr endfor %}")
    doc.save(str(template_path))

    assert _supports_fast_placeholder_render(str(template_path)) is False


def test_apilayer_template_mapping_adds_objective_pair(tmp_path):
    template_path = tmp_path / "apilayer_direct_template.docx"
    doc = Document()
    doc.add_paragraph("Legacy template objective")
    doc.add_paragraph("old.person@example.com")
    doc.save(str(template_path))

    employee = {
        "name": "Anouar ABDALLAH",
        "email": "user1@nextrh.com",
        "summary": "Results-driven security professional.",
        "experience": [],
        "education": [],
    }

    parsed_template = {
        "name": "Old Person",
        "email": "old.person@example.com",
        "objective": "Legacy template objective",
    }

    with patch("app.services.cv_generator_fallback.settings.APILAYER_API_KEY", "test-key"), patch(
        "app.services.cv_generator_fallback.parse_resume_from_file",
        return_value=parsed_template,
    ):
        pairs = _build_apilayer_template_pairs(str(template_path), employee)

    assert ("Legacy template objective", "Results-driven security professional.") in pairs
    assert ("old.person@example.com", "user1@nextrh.com") in pairs


def test_process_cv_uses_apilayer_template_mapping_for_unstructured_objective(tmp_path):
    template_path = tmp_path / "apilayer_process_template.docx"
    doc = Document()
    doc.add_paragraph("Legacy template objective")
    doc.add_paragraph("old.person@example.com")
    doc.save(str(template_path))

    employee = {
        "name": "Anouar ABDALLAH",
        "email": "user1@nextrh.com",
        "summary": "Results-driven security professional.",
        "experience": [],
        "education": [],
    }

    parsed_template = {
        "name": "Old Person",
        "email": "old.person@example.com",
        "objective": "Legacy template objective",
    }

    with patch("app.services.cv_generator_fallback.settings.APILAYER_API_KEY", "test-key"), patch(
        "app.services.cv_generator_fallback.settings.GROQ_API_KEY", ""
    ), patch(
        "app.services.cv_generator_fallback.parse_resume_from_file",
        return_value=parsed_template,
    ):
        out = process_cv(
            template_path=str(template_path),
            employee_data=employee,
            output_dir=str(tmp_path),
            output_pdf=False,
            language="original",
        )

    text, _paras = _extract_all_text(out)

    assert "Results-driven security professional." in text
    assert "Legacy template objective" not in text
    assert "user1@nextrh.com" in text


def test_template_execution_plan_recognizes_fast_placeholder_templates(tmp_path):
    template_path = tmp_path / "placeholder_template.docx"
    doc = Document()
    doc.add_paragraph("{{ full_name }}")
    doc.add_paragraph("{{email}}")
    doc.save(str(template_path))

    plan = _plan_template_execution(str(template_path))

    assert plan["template_mode"] == "placeholder"
    assert plan["use_fast_placeholder_render"] is True
    assert plan["strategy"] == "placeholder_fast"
    assert plan["run_body_section_replacement"] is False
    assert plan["run_textbox_section_replacement"] is False


def test_template_execution_plan_recognizes_structural_placeholder_templates(tmp_path):
    template_path = tmp_path / "structural_placeholder_template.docx"
    doc = Document()
    doc.add_paragraph("{{ full_name }}")
    doc.add_paragraph("{%tr for exp in experience %}")
    doc.add_paragraph("{{ exp.title }}")
    doc.add_paragraph("{%tr endfor %}")
    doc.save(str(template_path))

    plan = _plan_template_execution(str(template_path))

    assert plan["template_mode"] == "placeholder"
    assert plan["use_fast_placeholder_render"] is False
    assert plan["strategy"] == "placeholder_direct"
    assert plan["run_body_section_replacement"] is True
    assert plan["run_textbox_section_replacement"] is True


def test_template_execution_plan_recognizes_direct_template_families():
    rania_plan = _plan_template_execution(
        os.path.join(os.path.dirname(__file__), '..', 'test_templates', 'Resume-RaniaAmmar (1).docx')
    )
    canadian_plan = _plan_template_execution(
        os.path.join(os.path.dirname(__file__), '..', 'test_templates', '124-modele-cv-canadien.docx')
    )

    assert rania_plan["template_mode"] == "direct"
    assert rania_plan["strategy"] == "direct_body_sections"
    assert rania_plan["run_body_section_replacement"] is True

    assert canadian_plan["template_mode"] == "direct"
    assert canadian_plan["strategy"] == "direct_textbox_sections"
    assert canadian_plan["run_body_section_replacement"] is False
    assert canadian_plan["run_textbox_section_replacement"] is True


@patch('app.services.cv_generator_fallback._render_placeholder_template')
@patch('app.services.cv_generator_fallback._generate_with_replacement')
def test_process_cv_falls_back_to_direct_pipeline_for_structural_docxtpl_tags(
    mock_generate_with_replacement,
    mock_render_placeholder,
    tmp_path,
):
    template_path = tmp_path / "structural_placeholder_template.docx"
    doc = Document()
    doc.add_paragraph("{{ full_name }}")
    doc.add_paragraph("{%tr for exp in experience %}")
    doc.add_paragraph("{{ exp.title }}")
    doc.add_paragraph("{%tr endfor %}")
    doc.save(str(template_path))

    def _write_output(_template_path, _employee_data, output_path, *_args, **_kwargs):
        out_doc = Document()
        out_doc.add_paragraph("Rendered by direct pipeline")
        out_doc.save(output_path)
        return output_path

    mock_generate_with_replacement.side_effect = _write_output

    out = process_cv(
        template_path=str(template_path),
        employee_data={"name": "Jazil Gafsi"},
        output_dir=str(tmp_path),
        output_pdf=False,
    )

    assert os.path.exists(out)
    mock_render_placeholder.assert_not_called()
    mock_generate_with_replacement.assert_called_once()


@patch('app.services.cv_generator_fallback._generate_with_replacement')
@patch('app.services.cv_generator_fallback._render_placeholder_template')
def test_process_cv_retries_direct_pipeline_when_placeholder_render_raises_template_error(
    mock_render_placeholder,
    mock_generate_with_replacement,
    tmp_path,
):
    template_path = tmp_path / "simple_placeholder_template.docx"
    doc = Document()
    doc.add_paragraph("{{ full_name }}")
    doc.save(str(template_path))

    mock_render_placeholder.side_effect = TemplateSyntaxError("Encountered unknown tag 'tr'.", 12)

    def _write_output(_template_path, _employee_data, output_path, *_args, **_kwargs):
        out_doc = Document()
        out_doc.add_paragraph("Rendered by fallback after template error")
        out_doc.save(output_path)
        return output_path

    mock_generate_with_replacement.side_effect = _write_output

    out = process_cv(
        template_path=str(template_path),
        employee_data={"name": "Jazil Gafsi"},
        output_dir=str(tmp_path),
        output_pdf=False,
    )

    assert os.path.exists(out)
    mock_render_placeholder.assert_called_once()
    mock_generate_with_replacement.assert_called_once()


@patch('app.services.cv_generator_fallback._ai_get_replacements')
def test_direct_pipeline_prefers_deterministic_jinja_scalar_replacements_over_ai_guess(
    mock_ai_get_replacements,
    tmp_path,
):
    template_path = tmp_path / "direct_jinja_template.docx"
    doc = Document()
    doc.add_paragraph("{{ full_name }}")
    doc.add_paragraph("{{email}} • {{phone}}")
    doc.add_paragraph("{{ current_position }}")
    doc.add_paragraph("{%tr for exp in experience %}")
    doc.add_paragraph("{{ exp.title }}")
    doc.add_paragraph("{%tr endfor %}")
    doc.save(str(template_path))

    mock_ai_get_replacements.return_value = (
        [
            ("{{ current_position }}", "Project Manager"),
            ("{{ full_name }}", "Wrong Name"),
        ],
        [],
    )

    employee = {
        "name": "Jazil Gafsi",
        "email": "user3@nextrh.com",
        "phone": "+216 11 222 333",
        "title": "Chef de projet",
        "summary": "Professionnel experimente en gestion de projets.",
        "experience": [{"title": "Chef de projet", "company": "NextRH", "dates": "2020 - Present"}],
    }

    out = process_cv(
        template_path=str(template_path),
        employee_data=employee,
        output_dir=str(tmp_path),
        output_pdf=False,
        language="fr",
    )

    text, _paras = _extract_all_text(out)
    assert "Jazil Gafsi" in text
    assert "Chef de projet" in text
    assert "Project Manager" not in text
    assert "user3@nextrh.com" in text
    assert "+216 11 222 333" in text
    assert "{{" not in text


def test_direct_pipeline_preserves_header_line_breaks_for_scalar_jinja_placeholders(tmp_path):
    template_path = tmp_path / "linebreak_header_template.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("{{ full_name }}")
    p.add_run().add_break()
    p.add_run().add_break()
    p.add_run("{{email}} • {{phone}}")
    doc.add_paragraph("{%tr for exp in experience %}")
    doc.add_paragraph("{{ exp.title }}")
    doc.add_paragraph("{%tr endfor %}")
    doc.save(str(template_path))

    out = process_cv(
        template_path=str(template_path),
        employee_data={
            "name": "Jazil Gafsi",
            "email": "user3@nextrh.com",
            "phone": "+216 11 222 333",
            "experience": [{"title": "Chef de projet", "company": "NextRH", "dates": "2020 - Present"}],
        },
        output_dir=str(tmp_path),
        output_pdf=False,
    )

    first_para = Document(out).paragraphs[0].text
    assert "Jazil Gafsi" in first_para
    assert "user3@nextrh.com" in first_para
    assert "\n" in first_para


def test_resume_rania_template_rebuilds_docxtpl_scaffold_tables_as_paragraphs(tmp_path):
    with patch("app.services.cv_generator_fallback.settings.GROQ_API_KEY", ""):
        out = process_cv(
            template_path=os.path.join(os.path.dirname(__file__), '..', 'test_templates', 'Resume-RaniaAmmar.docx'),
            employee_data={
                "name": "Jazil Gafsi",
                "email": "user3@nextrh.com",
                "phone": "+216 11 222 333",
                "experience": [
                    {"title": "Chef de projet", "company": "NextStep", "dates": "2018 - Present", "description": "Pilotage de projets."},
                    {"title": "IT Infrastructure Administrateur", "company": "S2I", "dates": "2006 - 2016", "description": "Administration infrastructure."},
                ],
                "education": [
                    {"degree": "Technicien superieur en Reseau Informatique", "institution": "IMSET", "dates": "2005"},
                ],
                "skills": ["Serveurs", "Stockage", "Reseaux"],
                "certifications": [{"name": "Cert A", "date": "2020"}],
                "summary": "Professionnel experimente en gestion de projets.",
            },
            output_dir=str(tmp_path),
            output_pdf=False,
            language="fr",
        )

    result_doc = Document(out)
    paragraphs = [p.text.strip() for p in result_doc.paragraphs if p.text.strip()]
    normalized = [
        unicodedata.normalize("NFKD", p).encode("ascii", "ignore").decode("ascii")
        for p in paragraphs
    ]

    assert result_doc.tables == []
    assert "\n" in result_doc.paragraphs[0].text
    assert any("Technicien superieur en Reseau Informatique" in p for p in normalized)
    assert not any(p == "2005" for p in normalized)


def test_derived_title_is_not_injected_into_header_current_position(tmp_path):
    template_path = tmp_path / "derived_title_template.docx"
    doc = Document()
    doc.add_paragraph("{{ current_position }}")
    doc.add_paragraph("{%tr for exp in experience %}")
    doc.add_paragraph("{{ exp.title }}")
    doc.add_paragraph("{%tr endfor %}")
    doc.save(str(template_path))

    out = process_cv(
        template_path=str(template_path),
        employee_data={
            "name": "Jazil Gafsi",
            "experience": [{"title": "Chef de projet", "company": "NextRH", "dates": "2020 - Present"}],
        },
        output_dir=str(tmp_path),
        output_pdf=False,
        language="fr",
    )

    text = "\n".join(p.text.strip() for p in Document(out).paragraphs if p.text.strip())
    assert "Chef de projet" not in text
