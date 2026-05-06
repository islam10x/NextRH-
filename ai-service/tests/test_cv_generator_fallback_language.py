import os
import sys
from unittest.mock import patch


sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.services.cv_generator_fallback import (
    _AI_PROMPT,
    _infer_template_language,
    _language_instruction,
    _normalize_employee_language_for_generation,
    _normalize_language_code,
    _resolve_generation_language,
    process_cv,
)
from docx import Document


def test_language_code_normalization():
    assert _normalize_language_code("fr") == "fr"
    assert _normalize_language_code("fr-FR") == "fr"
    assert _normalize_language_code("EN-us") == "en"
    assert _normalize_language_code("original") == "original"


def test_language_instruction_for_french_is_strict():
    rule = _language_instruction("fr")
    assert "French" in rule
    assert "Do not output English" in rule


def test_ai_prompt_accepts_language_hint_placeholder():
    prompt = _AI_PROMPT.format(
        template_text="CV template sample",
        employee_json='{"name":"Test"}',
        populated_sections="summary/profile/objective",
        language_hint="Write strictly in French. Do not output English.",
        employee_full_name="Test User",
    )
    assert "Write strictly in French" in prompt


@patch('app.services.cv_generator_fallback.translate_cv_best_effort')
def test_employee_language_normalization_uses_target_language(mock_translate):
    source = {
        'summary': 'Ingenieur logiciel experimente',
        'experience': [{'title': 'Chef de projet', 'description': 'Pilotage des livraisons'}],
    }
    translated = {
        'summary': 'Experienced software engineer',
        'experience': [{'title': 'Project Manager', 'description': 'Delivery leadership'}],
    }
    mock_translate.return_value = translated

    result = _normalize_employee_language_for_generation(source, 'en')

    mock_translate.assert_called_once_with(source, 'en')
    assert result == translated


@patch('app.services.cv_generator_fallback.translate_cv_best_effort')
def test_employee_language_normalization_skips_original_language(mock_translate):
    source = {'summary': 'Resume deja dans la bonne langue'}

    result = _normalize_employee_language_for_generation(source, 'original')

    mock_translate.assert_not_called()
    assert result == source


def test_infer_template_language_detects_english(tmp_path):
    template_path = tmp_path / 'english_template.docx'
    doc = Document()
    doc.add_paragraph('Profile')
    doc.add_paragraph('Experience')
    doc.add_paragraph('Education')
    doc.add_paragraph('Skills')
    doc.save(str(template_path))

    assert _infer_template_language(str(template_path)) == 'en'
    assert _resolve_generation_language(str(template_path), 'original') == 'en'


def test_infer_template_language_detects_french(tmp_path):
    template_path = tmp_path / 'french_template.docx'
    doc = Document()
    doc.add_paragraph('Profil')
    doc.add_paragraph('Expérience professionnelle')
    doc.add_paragraph('Formation')
    doc.add_paragraph('Compétences')
    doc.save(str(template_path))

    assert _infer_template_language(str(template_path)) == 'fr'
    assert _resolve_generation_language(str(template_path), 'original') == 'fr'


@patch('app.services.cv_generator_fallback._normalize_employee_language_for_generation')
def test_process_cv_uses_inferred_template_language_for_original_mode(mock_normalize, tmp_path):
    template_path = tmp_path / 'english_process_template.docx'
    doc = Document()
    doc.add_paragraph('Profile', style='Heading 1')
    doc.add_paragraph('Template summary placeholder')
    doc.add_paragraph('Experience', style='Heading 1')
    doc.add_paragraph('Old experience placeholder')
    doc.save(str(template_path))

    mock_normalize.side_effect = lambda data, _lang: data

    process_cv(
        template_path=str(template_path),
        employee_data={
            'name': 'Jazil Gafsi',
            'summary': 'Ingenieur logiciel experimente.',
            'experience': [{'title': 'Chef de projet', 'company': 'NextStep', 'dates': '2018 - Present'}],
        },
        output_dir=str(tmp_path),
        output_pdf=False,
        language='original',
    )

    assert mock_normalize.call_args[0][1] == 'en'
