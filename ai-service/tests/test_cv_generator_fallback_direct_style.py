import os
import sys
from unittest.mock import patch

import pytest
from lxml import etree
from docx import Document
from docx.shared import Pt

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.services.cv_generator_fallback import (
    _build_replacements,
    _detect_personal_info,
    _extract_all_text,
    _identify_section_from_text,
    _identify_section_lxml,
    process_cv,
)

_TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), '..', 'test_templates')


def _fixture_missing(fname: str) -> bool:
    return not os.path.exists(os.path.join(_TEMPLATES_DIR, fname))


@pytest.mark.skip(
    reason="Requires unimplemented direct-template inline-entry rendering "
    "(mixed bold-title + normal run pattern within one paragraph)."
)
def test_direct_template_reuses_mixed_run_paragraph_pattern(tmp_path):
    template_path = tmp_path / "mixed_run_direct_template.docx"
    doc = Document()
    doc.add_paragraph("Experience", style="Heading 1")
    header = doc.add_paragraph(style="Normal")
    header.add_run("Senior Developer").bold = True
    header.add_run(", ACME  2020 - Present")
    body = doc.add_paragraph(style="Normal")
    body.paragraph_format.left_indent = Pt(9)
    body.add_run("Built internal tooling.")
    doc.add_paragraph("Education", style="Heading 1")
    edu = doc.add_paragraph(style="Normal")
    edu.add_run("BSc Computer Science").bold = True
    edu.add_run(" | MIT  2015")
    doc.save(str(template_path))

    out = process_cv(
        template_path=str(template_path),
        employee_data={
            "name": "Jazil Gafsi",
            "experience": [
                {
                    "title": "Chef de projet",
                    "company": "NextStep",
                    "dates": "2018 - Present",
                    "description": "Pilotage de projets.",
                }
            ],
            "education": [
                {
                    "degree": "Technicien superieur en Reseau Informatique",
                    "institution": "IMSET",
                    "dates": "2005",
                }
            ],
        },
        output_dir=str(tmp_path),
        output_pdf=False,
        language="fr",
    )

    result = Document(out)

    exp_para = next(p for p in result.paragraphs if "NextStep" in p.text)
    assert exp_para.text.startswith("Chef de projet, NextStep")
    assert exp_para.text.endswith("2018 - Present")
    assert exp_para.runs[0].text == "Chef de projet"
    assert exp_para.runs[0].bold is True
    assert any("NextStep" in run.text for run in exp_para.runs[1:])
    assert all(run.bold is not True for run in exp_para.runs[1:] if run.text)

    desc_para = next(p for p in result.paragraphs if "Pilotage de projets." in p.text)
    assert desc_para.paragraph_format.left_indent is not None
    assert round(desc_para.paragraph_format.left_indent.pt, 1) == 9.0

    edu_para = next(p for p in result.paragraphs if "IMSET" in p.text)
    assert edu_para.text.startswith("Technicien superieur en Reseau Informatique | IMSET")
    assert edu_para.text.endswith("2005")
    assert edu_para.runs[0].bold is True
    assert any("IMSET" in run.text for run in edu_para.runs[1:])


@pytest.mark.skipif(
    _fixture_missing('Resume-RaniaAmmar (1).docx'),
    reason="Fixture test_templates/Resume-RaniaAmmar (1).docx is not present in the repo.",
)
def test_resume_rania_direct_template_preserves_inline_entry_styles(tmp_path):
    template_path = os.path.join(
        os.path.dirname(__file__), '..', 'test_templates', 'Resume-RaniaAmmar (1).docx'
    )

    out = process_cv(
        template_path=template_path,
        employee_data={
            "name": "Jazil Gafsi",
            "email": "user3@nextrh.com",
            "phone": "+216 11 222 333",
            "experience": [
                {
                    "title": "Chef de projet",
                    "company": "NextStep",
                    "dates": "2018 - Present",
                    "description": "Pilotage de projets.",
                }
            ],
            "education": [
                {
                    "degree": "Technicien supérieur en Réseau Informatique",
                    "institution": "IMSET",
                    "dates": "2005",
                }
            ],
            "summary": "Professionnel expérimenté en gestion de projets.",
            "skills": ["Serveurs", "Stockage", "Réseaux"],
        },
        output_dir=str(tmp_path),
        output_pdf=False,
        language="fr",
    )

    result = Document(out)

    exp_para = next(p for p in result.paragraphs if "NextStep" in p.text)
    assert exp_para.text.startswith("Chef de projet, NextStep")
    assert exp_para.text.endswith("2018 - Present")
    assert exp_para.runs[0].bold is True
    assert all(run.bold is not True for run in exp_para.runs[1:] if run.text)

    desc_para = next(p for p in result.paragraphs if "Pilotage de projets." in p.text)
    assert desc_para.paragraph_format.left_indent is not None
    assert round(desc_para.paragraph_format.left_indent.pt, 1) == 9.0

    edu_para = next(p for p in result.paragraphs if "IMSET" in p.text)
    assert "Technicien supérieur en Réseau Informatique" in edu_para.text
    assert edu_para.runs[0].bold is True

    paragraphs = [p.text.strip() for p in result.paragraphs if p.text.strip()]
    assert "Tools & DevOps: Docker, Git, REST APIs, Postman, Linux." not in paragraphs
    assert paragraphs.count("Serveurs, Stockage, Réseaux") == 1


def test_label_value_content_row_is_not_misdetected_as_section_heading():
    doc = Document()
    para = doc.add_paragraph(style="Normal")
    para.add_run("Languages: ").bold = True
    para.add_run("French (Fluent), Arabic (Fluent), English (Fluent)")

    p_elem = etree.fromstring(para._p.xml.encode())

    assert _identify_section_lxml(p_elem) is None


def test_direct_template_preserves_paragraph_style_for_summary_block(tmp_path):
    template_path = tmp_path / "styled_summary_template.docx"
    doc = Document()
    doc.add_paragraph("Profile", style="Heading 1")
    summary_para = doc.add_paragraph(style="Intense Quote")
    summary_para.add_run("Template summary placeholder")
    doc.add_paragraph("Experience", style="Heading 1")
    doc.add_paragraph("Old experience content")
    doc.save(str(template_path))

    out = process_cv(
        template_path=str(template_path),
        employee_data={
            "name": "Jazil Gafsi",
            "summary": "Structured summary that should keep the template paragraph style.",
            "experience": [
                {
                    "title": "Chef de projet",
                    "company": "NextStep",
                    "dates": "2018 - Present",
                }
            ],
        },
        output_dir=str(tmp_path),
        output_pdf=False,
        language="original",
    )

    result = Document(out)
    summary_result = next(
        p for p in result.paragraphs
        if "Structured summary that should keep the template paragraph style." in p.text
    )

    assert summary_result.style.name == "Intense Quote"


def test_key_project_alias_maps_to_projects_section():
    assert _identify_section_from_text("Key Project") == "projects"
    assert _identify_section_from_text("Key Projects") == "projects"


@pytest.mark.skipif(
    _fixture_missing('124-modele-cv-canadien.docx'),
    reason="Fixture test_templates/124-modele-cv-canadien.docx is not present in the repo.",
)
def test_canadian_infographic_template_clears_split_company_textboxes(tmp_path):
    template_path = os.path.join(
        os.path.dirname(__file__), '..', 'test_templates', '124-modele-cv-canadien.docx'
    )

    with patch("app.services.cv_generator_fallback.settings.GROQ_API_KEY", ""):
        out = process_cv(
            template_path=template_path,
            employee_data={
                "name": "Jazil Gafsi",
                "email": "user3@nextrh.com",
                "phone": "",
                "experience": [
                    {
                        "title": "Chef de projet",
                        "company": "NextStep",
                        "dates": "2018 - Present",
                        "description": "Coordination des equipes en charge du projet.\nAssurer la relation client.\nPlanifier les projets et realiser des echeanciers.",
                    },
                    {
                        "title": "IT Infrastructure Administrateur",
                        "company": "S2I",
                        "dates": "2006 - 2016",
                        "description": "Gerer plusieurs projets commerciaux.\nVerifier la recevabilite des reclamations clients.",
                    },
                    {
                        "title": "IT Infrastructure Administrateur",
                        "company": "WAYCON",
                        "dates": "2016 - 2018",
                        "description": "Negocier, gerer les commandes et la prise de rendez-vous avec les clients.",
                    },
                ],
                "education": [
                    {
                        "degree": "Technicien superieur en Reseau Informatique",
                        "institution": "IMSET",
                        "dates": "2005",
                    }
                ],
                "skills": ["Serveurs", "Stockage", "Reseaux"],
                "languages": ["Francais", "Anglais"],
                "summary": "Professionnel experimente en gestion de projets.",
            },
            output_dir=str(tmp_path),
            output_pdf=False,
            language="original",
        )

    full_text, _paragraphs = _extract_all_text(out)

    assert "Jazil Gafsi" in full_text
    assert "NextStep" in full_text
    assert "Professionnel experimente en gestion de projets." in full_text
    assert "Serveurs, Stockage, Reseaux" in full_text
    assert "Francais, Anglais" in full_text
    assert "DESJARDINS" not in full_text
    assert "SOBEYS" not in full_text
    assert "CHU DE MONTRÉAL" not in full_text
    assert "LUCAS LEBLANC" not in full_text
    assert "Travail d’équipe" not in full_text
    assert "Espagnol" not in full_text
    assert "Allemand" not in full_text
    assert "Lycée Blaise Pascal" not in full_text
    assert "Université La Sorbonne" not in full_text
    assert "Chef de projet NextStep  —  2018 - Present" not in full_text


def test_infographic_contact_title_is_not_reused_as_name_replacement():
    paragraphs = [
        'Principales missions :',
        'Coordination des équipes en charge du projet.',
        'Assurer la relation client.',
        'Planifier les projets et réaliser des échéanciers.',
        'SOBEYS',
        'CHARGÉ DE PROJETS TI',
        'CHARGÉ DE PROJETS TI',
        'LUCAS LEBLANC',
        'lucas.leblanc @courriel.ca',
        '555-555-5555',
        'linkedin.com/lucas-leblanc',
        'Montréal, Canada',
    ]
    detected = _detect_personal_info('\n'.join(paragraphs), paragraphs)

    assert detected.get('name') == 'LUCAS LEBLANC'
    assert detected.get('title') == 'CHARGÉ DE PROJETS TI'

    pairs = _build_replacements(
        detected,
        {
            'name': 'Jazil Gafsi',
            'title': 'Directeur de projet',
            '_title_derived': False,
            'email': 'user3@nextrh.com',
            'phone': '',
            'linkedin': 'linkedin.com/in/jazil-gafsi',
            'address': '',
        },
    )

    assert ('CHARGÉ DE PROJETS TI', 'Jazil Gafsi') not in pairs
    assert ('CHARGÉ DE PROJETS TI', 'Directeur de projet') in pairs


@pytest.mark.skipif(
    _fixture_missing('Document 5.docx'),
    reason="Fixture test_templates/Document 5.docx is not present in the repo.",
)
def test_document5_template_keeps_jazil_summary_uniform_and_sections_grouped(tmp_path):
    template_path = os.path.join(
        os.path.dirname(__file__), '..', 'test_templates', 'Document 5.docx'
    )

    with patch("app.services.cv_generator_fallback.settings.GROQ_API_KEY", ""):
        out = process_cv(
            template_path=template_path,
            employee_data={
                "name": "Jazil Gafsi",
                "email": "j.gafsi@nextstep-it.com",
                "phone": "+216 98 772 808",
                "address": "31, Rue 1er Juin, Mutuelle-ville, 1082, Tunis",
                "linkedin": "",
                "summary": "Professionnel experimente en gestion de projets et logiciels de gestion de projet.",
                "skills": ["Serveurs", "Stockage", "Reseaux"],
                "experience": [
                    {
                        "start_date": "Février 2018",
                        "company": "NextStep",
                        "title": "Chef de projet",
                        "description": "Pilotage de projets.",
                    },
                    {
                        "start_date": "2006 - 2016",
                        "company": "S2I",
                        "title": "IT Infrastructure Administrateur",
                        "description": "Administration infrastructure.",
                    },
                ],
                "education": [
                    {
                        "end_date": "2005",
                        "institution": "IMSET.",
                        "degree": "Technicien supérieur en Réseau Informatique",
                    }
                ],
            },
            output_dir=str(tmp_path),
            output_pdf=False,
            language="fr",
        )

    result = Document(out)
    body_paragraphs = [p.text.strip() for p in result.paragraphs if p.text.strip()]

    assert "Expérience" in body_paragraphs
    assert "Compétences" in body_paragraphs
    assert "Formation" in body_paragraphs
    assert "Chef de projet" in body_paragraphs
    assert "NextStep  —  2018" in body_paragraphs
    assert "• Pilotage de projets." in body_paragraphs
    assert "Serveurs, Stockage, Reseaux" in body_paragraphs
    assert not any("Jeanne" in paragraph for paragraph in body_paragraphs)
    assert not any("Généraliste des ressources humaines" in paragraph for paragraph in body_paragraphs)
    assert not any("Stage de ressources humaines" in paragraph for paragraph in body_paragraphs)

    summary_para = next(
        p for p in result.paragraphs
        if "logiciels de gestion de projet." in p.text
    )
    non_empty_runs = [run for run in summary_para.runs if run.text.strip()]

    assert non_empty_runs
    first_style = (
        non_empty_runs[0].font.name,
        non_empty_runs[0].font.size,
        non_empty_runs[0].bold,
        non_empty_runs[0].italic,
    )
    assert all(
        (
            run.font.name,
            run.font.size,
            run.bold,
            run.italic,
        ) == first_style
        for run in non_empty_runs
    )