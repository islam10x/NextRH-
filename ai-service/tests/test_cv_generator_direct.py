"""Regression tests for direct generator mode on placeholder-free templates."""

from pathlib import Path

import pytest
from docx import Document

import app.services.cv_generator as cv_generator_module
from app.services.cv_generator import _detect_personal_info, _extract_all_text, process_cv


@pytest.fixture
def employee_data() -> dict:
    return {
        "name": "Aya BEN JEMAA",
        "email": "aya.benjemaa@nextstep.tn",
        "phone": "+216 98 772 817",
        "title": "Software Engineer",
        "linkedin": "linkedin.com/in/ayabenjemaa",
        "address": "Tunis, Tunisia",
        "summary": "Experienced engineer focused on scalable applications.",
        "skills": ["Python", "React", "Docker"],
        "experience": [
            {
                "title": "Software Engineer",
                "company": "Next Step",
                "dates": "2023 - Present",
                "description": "Built internal platforms and automation flows.",
            }
        ],
        "education": [
            {
                "degree": "BSc in Software Engineering",
                "institution": "Horizon School of Digital Tech",
                "dates": "2023 - 2026",
            }
        ],
        "certifications": ["Advanced Linux"],
        "languages": ["French", "English"],
    }


@pytest.fixture
def direct_template(tmp_path: Path) -> Path:
    template_path = tmp_path / "direct-template.docx"
    doc = Document()
    doc.add_paragraph("Lucas Huitorel")
    doc.add_paragraph("lucas.huitorel@example.com")
    doc.add_paragraph("+33 6 12 34 56 78")
    doc.add_paragraph("Senior Consultant")
    doc.add_paragraph("Professional summary placeholder")
    doc.save(template_path)
    return template_path


def test_text_extraction(direct_template: Path):
    full_text, paragraphs = _extract_all_text(str(direct_template))

    assert "Lucas Huitorel" in full_text
    assert "Senior Consultant" in full_text
    assert any("lucas.huitorel@example.com" in paragraph for paragraph in paragraphs)


def test_personal_info_detection(direct_template: Path):
    full_text, paragraphs = _extract_all_text(str(direct_template))
    detected = _detect_personal_info(full_text, paragraphs)

    assert detected.get("email") == "lucas.huitorel@example.com"
    assert "+33 6 12 34 56 78" in str(detected.get("phone") or "")


def test_full_generation(direct_template: Path, employee_data: dict, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr("app.services.cv_generator.settings.GROQ_API_KEY", "")
    monkeypatch.setattr("app.services.cv_generator_fallback.settings.GROQ_API_KEY", "")

    result = process_cv(
        template_path=str(direct_template),
        employee_data=employee_data,
        output_dir=str(tmp_path),
        output_pdf=False,
    )

    assert Path(result).exists()

    output_text, _ = _extract_all_text(result)
    assert employee_data["email"] in output_text
    assert "Aya" in output_text
    assert employee_data["phone"] in output_text


def test_process_cv_retries_with_legacy_primary_when_shared_runtime_fails(
    direct_template: Path,
    employee_data: dict,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    legacy_docx = tmp_path / "legacy.docx"
    Document().save(legacy_docx)

    monkeypatch.setattr(
        cv_generator_module,
        "_fallback_process_cv",
        lambda **_kwargs: (_ for _ in ()).throw(RuntimeError("boom")),
        raising=False,
    )
    monkeypatch.setattr(
        cv_generator_module,
        "generate_cv_document",
        lambda **_kwargs: {"docx_path": str(legacy_docx), "pdf_path": None},
    )

    result = process_cv(
        template_path=str(direct_template),
        employee_data=employee_data,
        output_dir=str(tmp_path),
        output_pdf=False,
    )

    assert result == str(legacy_docx)
