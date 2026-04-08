import json
import logging
import os
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import fitz
from docx import Document
from PIL import Image

logger = logging.getLogger(__name__)

_INDEX_FILENAME = ".template_index.json"


@dataclass
class TemplateFingerprint:
    tokens: Set[str]
    table_cols: List[int]


def _normalize_text(value: str) -> str:
    if not value:
        return ""
    normalized = unicodedata.normalize("NFD", value)
    normalized = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    normalized = normalized.lower()
    normalized = re.sub(r"[^a-z0-9\\s]", " ", normalized)
    normalized = re.sub(r"\\s+", " ", normalized).strip()
    return normalized


def _tokenize(text: str) -> Set[str]:
    normalized = _normalize_text(text)
    tokens = {t for t in normalized.split() if len(t) >= 3}
    return tokens


def _extract_docx_fingerprint(docx_path: str) -> TemplateFingerprint:
    doc = Document(docx_path)
    texts: List[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text:
            texts.append(paragraph.text)
    table_cols: List[int] = []
    for table in doc.tables:
        if table.columns:
            table_cols.append(len(table.columns))
        # capture header row text if any
        if table.rows:
            header = table.rows[0]
            for cell in header.cells:
                if cell.text:
                    texts.append(cell.text)
    tokens = _tokenize(" ".join(texts))
    return TemplateFingerprint(tokens=tokens, table_cols=table_cols)


def _extract_pdf_fingerprint(pdf_path: str) -> TemplateFingerprint:
    texts: List[str] = []
    with fitz.open(pdf_path) as doc:
        for page in doc:
            text = page.get_text("text") or ""
            if text.strip():
                texts.append(text)
    combined = " ".join(texts)
    tokens = _tokenize(combined)
    if not tokens:
        ocr_text = _ocr_pdf_text(pdf_path)
        tokens = _tokenize(ocr_text)
    return TemplateFingerprint(tokens=tokens, table_cols=[])


def _ocr_pdf_text(pdf_path: str) -> str:
    try:
        import pytesseract
    except Exception:
        pytesseract = None
    try:
        import easyocr
    except Exception:
        easyocr = None

    with fitz.open(pdf_path) as doc:
        if doc.page_count == 0:
            return ""
        page = doc[0]
        pix = page.get_pixmap(dpi=200, alpha=False)
        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

    if pytesseract is not None:
        try:
            return pytesseract.image_to_string(image) or ""
        except Exception:
            pass

    if easyocr is not None:
        try:
            reader = easyocr.Reader(["en", "fr"], gpu=False)
            results = reader.readtext(image)
            return " ".join([text for _box, text, _conf in results if text])
        except Exception:
            pass

    return ""


def _jaccard(a: Set[str], b: Set[str]) -> float:
    if not a or not b:
        return 0.0
    inter = len(a & b)
    union = len(a | b)
    return inter / union if union else 0.0


def _layout_score(a_cols: List[int], b_cols: List[int]) -> float:
    if not a_cols or not b_cols:
        return 0.0
    matches = 0
    remaining = b_cols[:]
    for col in a_cols:
        if col in remaining:
            matches += 1
            remaining.remove(col)
    return matches / max(len(a_cols), len(b_cols))


def _load_index(index_path: Path) -> Dict[str, Dict[str, object]]:
    if not index_path.exists():
        return {}
    try:
        raw = json.loads(index_path.read_text(encoding="utf-8"))
        return raw.get("templates", {}) if isinstance(raw, dict) else {}
    except Exception:
        return {}


def _save_index(index_path: Path, templates: Dict[str, Dict[str, object]]) -> None:
    payload = {"version": 1, "templates": templates}
    index_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _fingerprint_to_dict(fp: TemplateFingerprint) -> Dict[str, object]:
    return {"tokens": sorted(fp.tokens), "table_cols": fp.table_cols}


def _dict_to_fingerprint(data: Dict[str, object]) -> TemplateFingerprint:
    tokens = set(data.get("tokens") or [])
    table_cols = [int(v) for v in (data.get("table_cols") or [])]
    return TemplateFingerprint(tokens=tokens, table_cols=table_cols)


def _compute_or_load_fingerprint(path: Path, index: Dict[str, Dict[str, object]]) -> TemplateFingerprint:
    key = str(path)
    mtime = int(path.stat().st_mtime)
    cached = index.get(key)
    if cached and cached.get("mtime") == mtime:
        return _dict_to_fingerprint(cached.get("fingerprint", {}))

    fp = _extract_docx_fingerprint(str(path))
    index[key] = {"mtime": mtime, "fingerprint": _fingerprint_to_dict(fp)}
    return fp


def find_closest_template(
    scanned_pdf_path: str,
    templates_dir: Optional[str] = None,
    min_score: float = 0.12,
    top_k: int = 3,
) -> Dict[str, object]:
    pdf_fp = _extract_pdf_fingerprint(scanned_pdf_path)
    if not pdf_fp.tokens:
        return None

    if templates_dir:
        base_dir = Path(templates_dir)
    else:
        base_dir = Path(scanned_pdf_path).parent

    index_path = base_dir / _INDEX_FILENAME
    index = _load_index(index_path)

    candidates = list(base_dir.glob("*.docx"))
    if not candidates:
        return {"best": None, "top": []}

    scored: List[Tuple[str, float]] = []
    for candidate in candidates:
        try:
            fp = _compute_or_load_fingerprint(candidate, index)
        except Exception as exc:
            logger.warning(f"Template fingerprint failed for {candidate}: {exc}")
            continue

        token_score = _jaccard(pdf_fp.tokens, fp.tokens)
        layout_score = _layout_score(pdf_fp.table_cols, fp.table_cols)
        score = token_score * 0.85 + layout_score * 0.15
        scored.append((str(candidate), score))

    _save_index(index_path, index)
    scored.sort(key=lambda entry: entry[1], reverse=True)

    top_matches = scored[: max(top_k, 0)]
    if top_matches:
        preview = ", ".join([f"{Path(p).name} ({s:.3f})" for p, s in top_matches])
        logger.info(f"Template retrieval top matches: {preview}")

    best = top_matches[0] if top_matches else None
    if best and best[1] >= min_score:
        logger.info(f"Template retrieval matched {best[0]} with score {best[1]:.3f}")
        return {"best": best, "top": top_matches}
    return {"best": None, "top": top_matches}
