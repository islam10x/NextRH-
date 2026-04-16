# -*- coding: utf-8 -*-
import logging
import unicodedata
from pathlib import Path
from typing import Optional

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

logger = logging.getLogger(__name__)


def _normalize_text(value: str) -> str:
    if not value:
        return ""
    normalized = unicodedata.normalize("NFD", value)
    normalized = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    normalized = normalized.lower()
    normalized = " ".join(normalized.split())
    return normalized


def _extract_raw_text(pdf_path: str) -> str:
    text_parts = []
    with fitz.open(pdf_path) as doc:
        for page in doc:
            text_parts.append(page.get_text("text") or "")
    return "\n".join(text_parts).strip()


def _looks_like_annexe9(text: str) -> bool:
    norm = _normalize_text(text)
    return (
        "annexe 9" in norm
        and "curriculum" in norm
        and "vitae" in norm
        and "projet" in norm
        and "client" in norm
        and "periode" in norm
    )


def _build_annexe9_docx(output_path: Path) -> None:
    doc = Document()

    # Unicode-friendly base font
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    # Title
    p = doc.add_paragraph("Annexe 9 : Modèle de curriculum vitae")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.runs[0]
    run.bold = True
    run.font.size = Pt(14)

    # Personal info table
    info_table = doc.add_table(rows=0, cols=2)
    info_table.autofit = True

    def add_info_row(label: str, value: str) -> None:
        row = info_table.add_row().cells
        row[0].text = label
        row[1].text = value

    add_info_row("Nom et Prénom :", "{{ full_name or \"N/A\" }}")
    add_info_row("Date de naissance :", "{{ birth_date or \"N/A\" }}")
    add_info_row("Situation familiale :", "{{ marital_status or \"N/A\" }}")
    add_info_row("Date de recrutement :", "{{ hire_date or \"N/A\" }}")
    add_info_row("Dernier diplôme obtenu :", "{{ last_degree or \"N/A\" }}")
    add_info_row("Année :", "{{ last_degree_year or \"N/A\" }}")
    add_info_row("Fonction à assurer dans la mission :", "{{ current_position or \"N/A\" }}")
    add_info_row("Profil et connaissances :", "{{ professional_summary or \"N/A\" }}")

    # Experience table
    doc.add_paragraph("")
    heading = doc.add_paragraph("Expérience professionnelle générale dans le domaine de l'informatique")
    heading.runs[0].bold = True

    exp_table = doc.add_table(rows=1, cols=4)
    exp_table.style = "Table Grid"
    exp_table.autofit = True
    hdr = exp_table.rows[0].cells
    hdr[0].text = "Période"
    hdr[1].text = "Projet"
    hdr[2].text = "Client"
    hdr[3].text = "Durée globale"

    row = exp_table.add_row().cells
    row[0].text = ""
    row[1].text = ""
    row[2].text = ""
    row[3].text = ""

    # Projects table
    doc.add_paragraph("")
    heading = doc.add_paragraph(
        "Expérience professionnelle générale dans les projets similaires à la mission objet de l'appel d'offres"
    )
    heading.runs[0].bold = True

    proj_table = doc.add_table(rows=1, cols=4)
    proj_table.style = "Table Grid"
    proj_table.autofit = True
    hdr = proj_table.rows[0].cells
    hdr[0].text = "Période"
    hdr[1].text = "Projet"
    hdr[2].text = "Client"
    hdr[3].text = "Délai global"

    row = proj_table.add_row().cells
    row[0].text = ""
    row[1].text = ""
    row[2].text = ""
    row[3].text = ""

    # Signature
    doc.add_paragraph("")
    add_sig = doc.add_paragraph("Signature de l'intéressé")
    add_sig.runs[0].bold = True

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def build_docx_template_from_scanned_pdf(pdf_path: str, output_dir: Optional[str] = None) -> Optional[str]:
    raw_text = _extract_raw_text(pdf_path)
    if not raw_text:
        return None
    if not _looks_like_annexe9(raw_text):
        return None

    output_root = Path(output_dir) if output_dir else Path(pdf_path).parent
    output_path = output_root / f"{Path(pdf_path).stem}_autotemplate.docx"
    _build_annexe9_docx(output_path)
    logger.info(f"Auto-generated DOCX template from scanned PDF: {output_path}")
    return str(output_path)
