import logging
import os
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

logger = logging.getLogger(__name__)
_EASYOCR_READERS: Dict[Tuple[str, ...], object] = {}
_LOG_RAW_TEXT = os.getenv("CV_TEMPLATE_LOG_RAW") == "1"


@dataclass
class OCRLine:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float


def _pdf_has_text(pdf_path: str, min_chars: int = 40) -> bool:
    try:
        doc = fitz.open(pdf_path)
    except Exception:
        return False
    try:
        for page in doc:
            text = (page.get_text("text") or "").strip()
            if len(text) >= min_chars:
                return True
    finally:
        doc.close()
    return False


def _pdf_has_full_page_image(pdf_path: str, coverage_threshold: float = 0.85) -> bool:
    try:
        doc = fitz.open(pdf_path)
    except Exception:
        return False
    try:
        for page in doc:
            page_area = page.rect.width * page.rect.height
            if page_area <= 0:
                continue
            blocks = page.get_text("dict").get("blocks", [])
            max_image_area = 0.0
            for block in blocks:
                if block.get("type") != 1:
                    continue
                x0, y0, x1, y1 = block.get("bbox", (0, 0, 0, 0))
                area = max(0.0, (x1 - x0)) * max(0.0, (y1 - y0))
                if area > max_image_area:
                    max_image_area = area
            if max_image_area / page_area >= coverage_threshold:
                return True
    finally:
        doc.close()
    return False


def _docx_has_text(docx_path: str, min_chars: int = 40) -> bool:
    try:
        doc = Document(docx_path)
    except Exception:
        return False
    total = 0
    for paragraph in doc.paragraphs:
        total += len(paragraph.text or "")
        if total >= min_chars:
            return True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    total += len(paragraph.text or "")
                    if total >= min_chars:
                        return True
    return False


def is_pdf_scanned(pdf_path: str) -> bool:
    # Treat as scanned if there's no reliable text layer OR the page is
    # dominated by a full-page image (common for scanned templates).
    if not _pdf_has_text(pdf_path):
        return True
    if _pdf_has_full_page_image(pdf_path):
        return True
    return False


def is_docx_scanned(docx_path: str) -> bool:
    return not _docx_has_text(docx_path)


def _convert_docx_to_pdf(docx_path: str, output_dir: str) -> str:
    cmd = [
        "soffice",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        output_dir,
        docx_path,
    ]
    completed = subprocess.run(cmd, capture_output=True, text=True)
    if completed.returncode != 0:
        logger.error(f"DOCX to PDF conversion failed: {completed.stderr}")
        raise RuntimeError("Failed to convert DOCX to PDF. Ensure LibreOffice is installed.")
    pdf_path = os.path.join(output_dir, f"{Path(docx_path).stem}.pdf")
    if not os.path.exists(pdf_path):
        raise RuntimeError("DOCX to PDF conversion did not produce output file")
    return pdf_path


def _cluster_positions(values: Iterable[float], tolerance: float) -> List[float]:
    clusters: List[float] = []
    counts: List[int] = []
    for value in sorted(values):
        if not clusters:
            clusters.append(value)
            counts.append(1)
            continue
        if abs(value - clusters[-1]) <= tolerance:
            counts[-1] += 1
            clusters[-1] = (clusters[-1] * (counts[-1] - 1) + value) / counts[-1]
        else:
            clusters.append(value)
            counts.append(1)
    return clusters


def _remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "nil")
        tbl_borders.append(elem)
    tbl_pr.append(tbl_borders)


def _ocr_lines_tesseract(image: Image.Image, lang: str) -> List[OCRLine]:
    try:
        import pytesseract
        from pytesseract import Output
    except Exception as exc:
        raise RuntimeError("pytesseract is not available") from exc

    try:
        pytesseract.get_tesseract_version()
    except Exception as exc:
        raise RuntimeError("Tesseract binary not found in PATH") from exc

    data = pytesseract.image_to_data(
        image,
        output_type=Output.DICT,
        lang=lang,
        config="--psm 6",
    )
    lines: Dict[Tuple[int, int, int], Dict[str, object]] = {}
    for i in range(len(data["text"])):
        text = (data["text"][i] or "").strip()
        if not text:
            continue
        try:
            conf = float(data.get("conf", ["0"])[i])
        except Exception:
            conf = 0
        if conf < 10:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        entry = lines.setdefault(
            key,
            {"words": [], "x0": float("inf"), "y0": float("inf"), "x1": 0.0, "y1": 0.0},
        )
        entry["words"].append(text)
        left = data["left"][i]
        top = data["top"][i]
        width = data["width"][i]
        height = data["height"][i]
        entry["x0"] = min(entry["x0"], left)
        entry["y0"] = min(entry["y0"], top)
        entry["x1"] = max(entry["x1"], left + width)
        entry["y1"] = max(entry["y1"], top + height)

    results: List[OCRLine] = []
    for entry in lines.values():
        words = entry["words"]
        text = " ".join(words).strip()
        if not text:
            continue
        results.append(
            OCRLine(
                text=text,
                x0=float(entry["x0"]),
                y0=float(entry["y0"]),
                x1=float(entry["x1"]),
                y1=float(entry["y1"]),
            )
        )
    return results


def _ocr_lines_easyocr(image: Image.Image, lang: str) -> List[OCRLine]:
    try:
        import easyocr
    except Exception as exc:
        raise RuntimeError("easyocr is not available") from exc

    lang_map = {"eng": "en", "fra": "fr", "ara": "ar"}
    langs = []
    for token in lang.split("+"):
        token = token.strip().lower()
        if token in lang_map:
            langs.append(lang_map[token])
    if not langs:
        langs = ["en"]

    key = tuple(langs)
    reader = _EASYOCR_READERS.get(key)
    if reader is None:
        reader = easyocr.Reader(langs, gpu=False)
        _EASYOCR_READERS[key] = reader
    results = reader.readtext(image)
    lines: List[OCRLine] = []
    for box, text, conf in results:
        if not text or conf < 0.2:
            continue
        xs = [p[0] for p in box]
        ys = [p[1] for p in box]
        lines.append(
            OCRLine(
                text=text.strip(),
                x0=min(xs),
                y0=min(ys),
                x1=max(xs),
                y1=max(ys),
            )
        )
    return lines


def _ocr_lines(image: Image.Image, lang: str) -> List[OCRLine]:
    try:
        return _ocr_lines_tesseract(image, lang)
    except Exception as tesseract_err:
        logger.warning(f"Tesseract OCR failed, trying EasyOCR: {tesseract_err}")
        return _ocr_lines_easyocr(image, lang)


def _lines_from_text_layer(words: List[Tuple[float, float, float, float, str, int, int, int]]) -> List[OCRLine]:
    lines: Dict[Tuple[int, int], Dict[str, object]] = {}
    for x0, y0, x1, y1, text, block_no, line_no, _word_no in words:
        clean = (text or "").strip()
        if not clean:
            continue
        key = (block_no, line_no)
        entry = lines.setdefault(
            key,
            {"words": [], "x0": float("inf"), "y0": float("inf"), "x1": 0.0, "y1": 0.0},
        )
        entry["words"].append(clean)
        entry["x0"] = min(entry["x0"], x0)
        entry["y0"] = min(entry["y0"], y0)
        entry["x1"] = max(entry["x1"], x1)
        entry["y1"] = max(entry["y1"], y1)

    results: List[OCRLine] = []
    for entry in lines.values():
        text = " ".join(entry["words"]).strip()
        if not text:
            continue
        results.append(
            OCRLine(
                text=text,
                x0=float(entry["x0"]),
                y0=float(entry["y0"]),
                x1=float(entry["x1"]),
                y1=float(entry["y1"]),
            )
        )
    return results


def _group_lines_into_rows(lines: List[OCRLine], tolerance: float) -> List[List[OCRLine]]:
    if not lines:
        return []
    rows: List[List[OCRLine]] = []
    current: List[OCRLine] = []
    current_y = None
    for line in sorted(lines, key=lambda l: (l.y0, l.x0)):
        if current_y is None:
            current = [line]
            current_y = line.y0
            continue
        if abs(line.y0 - current_y) <= tolerance:
            current.append(line)
            current_y = (current_y + line.y0) / 2
        else:
            rows.append(sorted(current, key=lambda l: l.x0))
            current = [line]
            current_y = line.y0
    if current:
        rows.append(sorted(current, key=lambda l: l.x0))
    return rows


def _build_docx_from_lines(
    doc: Document,
    page_lines: List[OCRLine],
    page_width_px: int,
    page_height_px: int,
    page_width_in: float,
    page_height_in: float,
) -> None:
    if not page_lines:
        doc.add_paragraph("")
        return

    avg_height = sum((line.y1 - line.y0) for line in page_lines) / max(len(page_lines), 1)
    row_tolerance = max(6.0, avg_height * 0.6)
    rows = _group_lines_into_rows(page_lines, tolerance=row_tolerance)

    x_positions = [line.x0 for line in page_lines]
    column_tolerance = max(24.0, page_width_px * 0.04)
    column_centers = _cluster_positions(x_positions, column_tolerance)
    use_columns = len(column_centers) >= 2

    if use_columns:
        table = doc.add_table(rows=0, cols=len(column_centers))
        table.autofit = True
        _remove_table_borders(table)
        for row_lines in rows:
            row = table.add_row().cells
            col_text: Dict[int, List[str]] = {i: [] for i in range(len(column_centers))}
            for line in row_lines:
                col_index = min(
                    range(len(column_centers)),
                    key=lambda i: abs(line.x0 - column_centers[i]),
                )
                col_text[col_index].append(line.text)
            for idx, texts in col_text.items():
                row[idx].text = " ".join(texts).strip()
    else:
        last_y = None
        for row_lines in rows:
            text = " ".join(line.text for line in row_lines).strip()
            if not text:
                continue
            y0 = min(line.y0 for line in row_lines)
            if last_y is not None and (y0 - last_y) > avg_height * 1.6:
                doc.add_paragraph("")
            paragraph = doc.add_paragraph(text)
            min_x = min(line.x0 for line in row_lines)
            indent_inches = (min_x / max(page_width_px, 1)) * page_width_in
            paragraph.paragraph_format.left_indent = Inches(indent_inches)
            last_y = y0


def rebuild_scanned_pdf_to_docx(
    pdf_path: str,
    output_dir: Optional[str] = None,
    lang: str = "eng+fra",
    dpi: int = 300,
) -> str:
    pdf_path = str(Path(pdf_path))
    if output_dir is None:
        output_dir = str(Path(pdf_path).parent)

    base_name = Path(pdf_path).stem
    rebuilt_path = Path(output_dir) / f"{base_name}_rebuilt.docx"

    doc = Document()
    with fitz.open(pdf_path) as pdf:
        for page_index, page in enumerate(pdf):
            if _LOG_RAW_TEXT:
                raw_text = (page.get_text("text") or "").strip()
                if raw_text:
                    logger.info(f"[TEMPLATE] Raw text page {page_index + 1}:\n{raw_text}")

            words = page.get_text("words") or []
            page_lines: List[OCRLine] = []
            if len(words) >= 10:
                page_lines = _lines_from_text_layer(words)
                if _LOG_RAW_TEXT:
                    logger.info(
                        f"[TEMPLATE] Using text-layer lines: {len(page_lines)} lines on page {page_index + 1}"
                    )
            if not page_lines:
                pix = page.get_pixmap(dpi=dpi, alpha=False)
                image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                page_lines = _ocr_lines(image, lang)
                if _LOG_RAW_TEXT:
                    logger.info(
                        f"[TEMPLATE] Using OCR lines: {len(page_lines)} lines on page {page_index + 1}"
                    )

            page_width_in = page.rect.width / 72.0
            page_height_in = page.rect.height / 72.0

            if page_index == 0:
                section = doc.sections[0]
                section.page_width = Inches(page_width_in)
                section.page_height = Inches(page_height_in)

            _build_docx_from_lines(
                doc,
                page_lines,
                page_width_px=int(page.rect.width),
                page_height_px=int(page.rect.height),
                page_width_in=page_width_in,
                page_height_in=page_height_in,
            )

            if page_index < len(pdf) - 1:
                doc.add_page_break()

    doc.save(rebuilt_path)
    logger.info(f"OCR rebuilt DOCX saved at: {rebuilt_path}")
    return str(rebuilt_path)


def rebuild_scanned_template(
    template_path: str,
    output_dir: Optional[str] = None,
    lang: str = "eng+fra",
) -> str:
    template_path = str(Path(template_path))
    ext = Path(template_path).suffix.lower()
    if ext == ".pdf":
        return rebuild_scanned_pdf_to_docx(template_path, output_dir=output_dir, lang=lang)
    if ext == ".docx":
        if output_dir is None:
            output_dir = str(Path(template_path).parent)
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_pdf = _convert_docx_to_pdf(template_path, temp_dir)
            return rebuild_scanned_pdf_to_docx(temp_pdf, output_dir=output_dir, lang=lang)
    raise ValueError(f"Unsupported template type for OCR rebuild: {ext}")
