import docx
from .base_parser import BaseParser
from typing import Any, Dict

class DocxParser(BaseParser):
    def parse(self, file_path: str) -> str:
        """
        Extract text from a DOCX file using python-docx.
        """
        text = ""
        try:
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            text = '\n'.join(full_text)
        except Exception as e:
             raise ValueError(f"Failed to parse DOCX: {e}")
        return text

    def extract_metadata(self, file_path: str) -> Dict[str, Any]:
        """
        Extract metadata from a DOCX file.
        """
        metadata = {}
        try:
            doc = docx.Document(file_path)
            core_props = doc.core_properties
            metadata = {
                "author": core_props.author,
                "created": core_props.created,
                "modified": core_props.modified,
                "title": core_props.title,
                "subject": core_props.subject
            }
        except Exception as e:
             # metadata extraction failure shouldn't necessarily block parsing
             pass
        return metadata
