import docx
import zipfile

docx_path = "EVAL_NAME_JOHN_EVAL_NAME_DOE_CV.docx"
doc = docx.Document(docx_path)
full_text = []
for para in doc.paragraphs:
    full_text.append(para.text)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                full_text.append(para.text)
                
print("=== PYTHON-DOCX EXTRACTED TEXT ===")
print("\n".join(t for t in full_text if t.strip()).encode("utf-8", errors="ignore").decode("utf-8", errors="ignore"))

try:
    with zipfile.ZipFile(docx_path, 'r') as zf:
        xml_content = ""
        for item in zf.namelist():
            if item.endswith(".xml") and "word/" in item:
                xml_content += zf.read(item).decode('utf-8')
        
        print("\n=== XML RAW SEARCH ===")
        print("PYTHON found:", "PYTHON" in xml_content)
        print("REACT found:", "REACT" in xml_content)
        print("AWS found:", "AWS" in xml_content)
        print("SUMMARY found:", "SUMMARY" in xml_content)
except Exception as e:
    print(e)
