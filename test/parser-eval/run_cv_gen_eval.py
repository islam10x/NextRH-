import os
import time
import json
import requests
import docx
from pathlib import Path

# Setup paths
BASE_DIR = Path(r"c:\Users\islam\Projects\NextRH-")
TEMPLATES_DIR = BASE_DIR / "test" / "test-templates"
API_URL = "http://127.0.0.1:8000/api/v1/generation/cv"

# Dummy profile data with EVAL tags
DUMMY_PROFILE = {
    "name": "EVAL_NAME_JOHN EVAL_NAME_DOE",
    "email": "eval_email@evaldomain.com",
    "phone": "999-EVAL-PHONE",
    "title": "EVAL_TITLE_ENGINEER",
    "address": "EVAL_ADDRESS_CITY",
    "linkedin": "linkedin.com/in/eval_linkedin",
    "summary": "EVAL_SUMMARY_PROFESSIONAL",
    "skills": ["EVAL_SKILL_PYTHON", "EVAL_SKILL_REACT"],
    "experience": [
        {
            "title": "EVAL_EXP_ROLE",
            "company": "EVAL_EXP_COMP",
            "dates": "2020 - 2023",
            "description": "EVAL_EXP_DESC"
        }
    ],
    "education": [
        {
            "degree": "EVAL_EDU_DEG",
            "institution": "EVAL_EDU_UNI",
            "dates": "2015 - 2019"
        }
    ],
    "languages": ["EVAL_LANG_EN"],
    "certifications": ["EVAL_CERT_AWS"]
}

def extract_text_from_docx(docx_path):
    doc = docx.Document(docx_path)
    full_text = []
    # Extract text from paragraphs
    for para in doc.paragraphs:
        full_text.append(para.text)
    # Extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    full_text.append(para.text)
    
    # Text in shapes/textboxes is harder with python-docx, but python-docx 
    # doesn't natively support textboxes. However, for a basic eval it might suffice, 
    # or we can do raw XML parsing. Let's do a basic raw XML search if needed.
    
    # Actually, the best way to be sure we don't miss textboxes is to read the raw XML.
    import zipfile
    try:
        with zipfile.ZipFile(docx_path, 'r') as zf:
            xml_content = ""
            for item in zf.namelist():
                if item.endswith(".xml") and "word/" in item:
                    xml_content += zf.read(item).decode('utf-8')
            return xml_content
    except Exception as e:
        print(f"Failed to read XML from {docx_path}: {e}")
        return "\n".join(full_text)

def run_eval():
    templates = [f for f in os.listdir(TEMPLATES_DIR) if f.endswith(".docx")]
    print(f"Found {len(templates)} templates to evaluate.\n")
    
    metrics = {
        "contact_block": 0,
        "professional_summary": 0,
        "work_experience": 0,
        "education_records": 0,
        "skills_list": 0,
        "certifications_list": 0,
        "cache_hits": 0,
        "total_templates": len(templates)
    }

    if not templates:
        print("No templates found in", TEMPLATES_DIR)
        return

    for idx, template_name in enumerate(templates, 1):
        template_path = TEMPLATES_DIR / template_name
        print(f"[{idx}/{len(templates)}] Evaluating: {template_name}")
        
        # --- PASS 1: Generate CV (Cache Miss) ---
        start_time_pass1 = time.time()
        try:
            with open(template_path, "rb") as f:
                files = {"template": (template_name, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
                data = {
                    "employee_data": json.dumps(DUMMY_PROFILE),
                    "engine": "fallback",
                    "output_format": "docx"
                }
                response1 = requests.post(API_URL, files=files, data=data)
            pass1_duration = time.time() - start_time_pass1

            if response1.status_code != 200:
                print(f"  [ERROR] Pass 1 Failed: HTTP {response1.status_code} - {response1.text}")
                continue
            
            # Save the returned docx to check contents
            out_docx = f"output_{idx}.docx"
            with open(out_docx, "wb") as f_out:
                f_out.write(response1.content)

        except Exception as e:
            print(f"  [ERROR] Request Failed: {e}")
            continue

        # --- Check Field Population ---
        doc_xml = extract_text_from_docx(out_docx)
        
        # Contact block is considered 100% if name, email, or phone is found (AI handles mapping)
        if "EVAL_NAME_JOHN" in doc_xml or "eval_email" in doc_xml or "999-EVAL-PHONE" in doc_xml:
            metrics["contact_block"] += 1
            
        if "EVAL_SUMMARY_PROFESSIONAL" in doc_xml:
            metrics["professional_summary"] += 1
            
        if "EVAL_EXP_ROLE" in doc_xml or "EVAL_EXP_COMP" in doc_xml:
            metrics["work_experience"] += 1
            
        if "EVAL_EDU_DEG" in doc_xml or "EVAL_EDU_UNI" in doc_xml:
            metrics["education_records"] += 1
            
        if "EVAL_SKILL_PYTHON" in doc_xml or "EVAL_SKILL_REACT" in doc_xml:
            metrics["skills_list"] += 1
            
        if "EVAL_CERT_AWS" in doc_xml:
            metrics["certifications_list"] += 1

        # --- PASS 2: Generate CV (Cache Hit check) ---
        # The caching relies on the template hash. So sending the exact same template
        # should trigger the mapping cache in cv_generator_fallback.py
        start_time_pass2 = time.time()
        try:
            with open(template_path, "rb") as f:
                files = {"template": (template_name, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
                data = {
                    "employee_data": json.dumps(DUMMY_PROFILE),
                    "engine": "fallback",
                    "output_format": "docx"
                }
                response2 = requests.post(API_URL, files=files, data=data)
            pass2_duration = time.time() - start_time_pass2
            
            # Simple heuristic: if the second pass is significantly faster, it hit the cache.
            # Usually fallback generation takes 5-15s, cache hit takes <3s.
            if response2.status_code == 200 and pass2_duration < pass1_duration * 0.7:
                metrics["cache_hits"] += 1
                print(f"  [HIT] Cache Hit (Pass 1: {pass1_duration:.1f}s, Pass 2: {pass2_duration:.1f}s)")
            else:
                print(f"  [MISS] Cache Miss or no speedup (Pass 1: {pass1_duration:.1f}s, Pass 2: {pass2_duration:.1f}s)")

        except Exception as e:
            print(f"  [ERROR] Pass 2 Failed: {e}")
            
        # Clean up output
        if os.path.exists(out_docx):
            os.remove(out_docx)

    # --- Print Summary ---
    print("\n==================================================")
    print("STANDARD CV GENERATION ENGINE — FIELD POPULATION ACCURACY")
    print(f"Total Templates: {metrics['total_templates']}")
    print("==================================================")
    
    def pct(val):
        return (val / metrics["total_templates"]) * 100 if metrics["total_templates"] > 0 else 0

    print(f"Contact block:         {pct(metrics['contact_block']):.1f}% (Target: 100%)")
    print(f"Professional summary:  {pct(metrics['professional_summary']):.1f}% (Target: 100%)")
    print(f"Work experience:       {pct(metrics['work_experience']):.1f}% (Target: 100%)")
    print(f"Education records:     {pct(metrics['education_records']):.1f}% (Target: 100%)")
    print(f"Skills list:           {pct(metrics['skills_list']):.1f}% (Target: 93%)")
    print(f"Certifications list:   {pct(metrics['certifications_list']):.1f}% (Target: 93%)")
    print(f"Cache hit rate:        {pct(metrics['cache_hits']):.1f}% (Target: 100%)")
    print("==================================================")

if __name__ == "__main__":
    run_eval()
